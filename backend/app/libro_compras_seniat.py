"""El Libro de Compras en el formato del SENIAT, como .xlsx.

Copia la planilla que usa el contador (columnas A a AP: detalle del
documento, totales, compras nacionales por alicuota, internacionales,
IGTF y retenciones, y el resumen de creditos fiscales al pie). Los montos van
en BOLIVARES: los congelados en cada factura al guardarla.

Aca si se usa openpyxl, a diferencia de los CSV de `exportar_csv`: este
formato tiene dos filas de encabezado combinadas, sumas y un resumen debajo,
y un CSV no lo puede dar. Lo que el ERP todavia no lleva (importaciones,
IGTF, retenciones de IVA) sale en cero, como en la planilla.
"""

import datetime
import io
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import impuestos, schemas

GRIS = PatternFill("solid", fgColor="D9D9D9")
NEGRITA = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONTO = "#,##0.00"
PORCENTAJE = "0.00%"

ENCABEZADOS = [
    "N° operación", "Fecha del documento", "RIF", "Nombre/Razón social", "Tipo de Documento",
    "N° de Factura", "N° Nota de Crédito", "N° Nota de Débito", "N° de control", "Tipo de transacción",
    "N° Factura afectada", "Total compras", "Total compras con IVA", "Total compras exentas",
    "Base imponible (16%)", "Alicuota (16%)", "IVA 16%", "Base imponible (8%)", "Alicuota (8%)", "IVA 8%",
    "Base imponible (31%)", "Alicuota (31%)", "IVA 31%",
    "Total compras", "Total compras con IVA", "Total compras exentas",
    "Número de Declaración Única de Aduana", "Número de expediente de Importación",
    "Valor total de las importaciones definitivas",
    "Base imponible (16%)", "Alicuota Int. (16%)", "IVA Int. 16%", "Base imponible (8%)",
    "Alicuota Int. (8%)", "IVA Int. 8%", "Base imponible (31%)", "Alicuota Int. (31%)", "IVA Int. 31%",
    "Igtf", "Fecha Retención", "N° Retención", "IVA retenido",
]
# Grupos de la fila de arriba: (columna desde, hasta, texto). 1 = A.
GRUPOS = [
    (1, 11, "DETALLE DEL DOCUMENTO"), (12, 14, "TOTALES"), (15, 23, "COMPRAS NACIONALES"),
    (24, 26, "TOTALES INTERNACIONALES"), (27, 38, "COMPRAS INTERNACIONALES"), (39, 39, "IGTF"),
    (40, 42, "RETENCIONES"),
]
# Columnas de montos: van en cero cuando no aplican, y llevan suma al pie.
MONTOS = [12, 13, 14, 15, 17, 18, 20, 21, 23, 24, 25, 26, 29, 30, 32, 33, 35, 36, 38, 39, 42]
# Columnas de alicuota y su porcentaje fijo.
ALICUOTAS = {16: 16, 19: 8, 22: 31, 31: 16, 34: 8, 37: 31}
# Donde cae cada alicuota nacional: (base, iva).
COLUMNAS_ALICUOTA = {16.0: (15, 17), 8.0: (18, 20), 31.0: (21, 23)}

FILA_GRUPOS = 6
FILA_ENCABEZADO = 7
PRIMERA_FILA = 8


def _f(x: datetime.date) -> str:
    return x.strftime("%d/%m/%Y")


def generar(libro: schemas.LibroCompras, fiscal, desde: datetime.date, hasta: datetime.date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Libro de Compras"

    rif = impuestos.normalizar_rif(fiscal.rif or "")
    razon = (fiscal.razon_social or "").strip()
    ws["C1"] = " - ".join(x for x in (razon, rif) if x) or "Razón social y RIF: cárgalos en Impuestos"
    ws["C2"] = f"Direccion:  {fiscal.direccion}" if fiscal.direccion else "Direccion:"
    ws["C3"] = "Libro de Compras"
    ws["C4"] = f"Desde {_f(desde)} Hasta {_f(hasta)}"
    for fila in range(1, 5):
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=13)
        ws.cell(fila, 3).font = NEGRITA

    for desde_col, hasta_col, texto in GRUPOS:
        if hasta_col > desde_col:
            ws.merge_cells(start_row=FILA_GRUPOS, start_column=desde_col, end_row=FILA_GRUPOS, end_column=hasta_col)
        for col in range(desde_col, hasta_col + 1):
            ws.cell(FILA_GRUPOS, col).fill = GRIS
        c = ws.cell(FILA_GRUPOS, desde_col, texto)
        c.font, c.alignment = NEGRITA, CENTRO
    for col, texto in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(FILA_ENCABEZADO, col, texto)
        c.font, c.fill, c.alignment = NEGRITA, GRIS, CENTRO

    alicuota = float(libro.tasa_iva)
    base_col, iva_col = COLUMNAS_ALICUOTA.get(alicuota, COLUMNAS_ALICUOTA[16.0])
    filas: List[schemas.FilaLibroCompras] = libro.filas
    r = PRIMERA_FILA
    for n, f in enumerate(filas, start=1):
        valores = {
            1: n,
            2: _f(f.fecha_emision),
            3: impuestos.normalizar_rif(f.proveedor_rif or ""),
            4: f.proveedor_nombre,
            5: f.tipo,
            6: f.numero_factura if f.tipo == "FAC" else "",
            7: f.numero_nota if f.tipo == "NC" else "",
            8: "",
            9: f.numero_control,
            10: "01-REG",
            11: f.factura_afectada or "--",
            27: "-",
            28: "--",
        }
        for col, v in valores.items():
            ws.cell(r, col, v)
        # La planilla quiere los montos que no aplican en cero, no vacios.
        for col in MONTOS:
            ws.cell(r, col, 0).number_format = MONTO
        for col, pct in ALICUOTAS.items():
            ws.cell(r, col, pct / 100).number_format = PORCENTAJE
        if f.total_bs is not None:
            gravado, exento, iva = f.gravado_bs or 0, f.exento_bs or 0, f.iva_bs or 0
            ws.cell(r, 12, f.total_bs)
            ws.cell(r, 13, round(gravado + iva, 2))
            ws.cell(r, 14, exento)
            ws.cell(r, base_col, gravado)
            ws.cell(r, iva_col, iva)
        else:
            # Sin tasa no hay Bs: un cero pareceria una compra de nada.
            for col in (12, 13, 14, base_col, iva_col):
                ws.cell(r, col).value = None
            ws.cell(r, 4, f"{f.proveedor_nombre} (SIN TASA: falta la tasa de su fecha)")
        r += 1

    ultima = r - 1
    for col in MONTOS:
        letra = get_column_letter(col)
        c = ws.cell(r, col, f"=SUM({letra}{PRIMERA_FILA}:{letra}{ultima})" if filas else 0)
        c.number_format, c.font = MONTO, NEGRITA

    _resumen(ws, r + 2, filas, alicuota)

    anchos = {1: 10, 2: 12, 3: 13, 4: 34, 5: 10, 6: 14, 7: 12, 8: 12, 9: 14, 10: 12, 11: 14}
    for col in range(1, len(ENCABEZADOS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = anchos.get(col, 15)
    ws.row_dimensions[FILA_ENCABEZADO].height = 45
    ws.freeze_panes = ws.cell(PRIMERA_FILA, 5)

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _resumen(ws, r0: int, filas: List[schemas.FilaLibroCompras], alicuota: float) -> None:
    """El cuadro de creditos fiscales del pie, con facturas y notas de
    credito por separado, como la planilla."""

    def total(tipo: str, campo) -> float:
        if campo is None:
            return 0.0
        return round(sum(getattr(f, campo) or 0 for f in filas if f.tipo == tipo), 2)

    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=2)
    for col, texto in ((1, "Resumen"), (3, "Facturas/Notas de Débito"), (5, "Notas de Crédito"), (7, "Total Neto")):
        if col > 1:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r0, end_column=col + 1)
        c = ws.cell(r0, col, texto)
        c.font, c.fill, c.alignment = NEGRITA, GRIS, CENTRO
    subtitulos = ["Créditos Fiscales", "Base Imponible", "Crédito Fiscal", "Base Imponible",
                  "Crédito Fiscal", "Base Imponible", "Crédito Fiscal"]
    for col, texto in enumerate(subtitulos, start=2):
        c = ws.cell(r0 + 1, col, texto)
        c.font, c.fill, c.alignment = NEGRITA, GRIS, CENTRO

    def gravadas(pct: float):
        return ("gravado_bs", "iva_bs") if pct == alicuota else (None, None)

    conceptos = [
        ("Compras Internas no Gravadas", ("exento_bs", None)),
        ("Importaciones Gravadas por Alícuota Reducida", (None, None)),
        ("Importaciones Gravadas por Alícuota General", (None, None)),
        ("Importaciones Gravadas por Alícuota General más Adicional", (None, None)),
        ("Compras Internas Gravadas sólo por Alícuota General", gravadas(16.0)),
        ("Compras Internas Gravadas por Alícuota General más Adicional", gravadas(31.0)),
        ("Compras Internas Gravadas por Alícuota Reducida", gravadas(8.0)),
        ("Ajustes a los Créditos Fiscales de Periodos Anteriores", (None, None)),
    ]
    r = r0 + 2
    for n, (texto, (base, credito)) in enumerate(conceptos, start=1):
        ws.cell(r, 1, n)
        ws.cell(r, 2, texto)
        montos = (total("FAC", base), total("FAC", credito), total("NC", base), total("NC", credito))
        for col, v in zip(range(3, 7), montos):
            ws.cell(r, col, v).number_format = MONTO
        ws.cell(r, 7, f"=C{r}+E{r}").number_format = MONTO
        ws.cell(r, 8, f"=D{r}+F{r}").number_format = MONTO
        r += 1
    ws.cell(r, 1, len(conceptos) + 1)
    ws.cell(r, 2, "Total Compras y Créditos Fiscales del Periodo").font = NEGRITA
    for col in range(3, 9):
        letra = get_column_letter(col)
        c = ws.cell(r, col, f"=SUM({letra}{r0 + 2}:{letra}{r - 1})")
        c.number_format, c.font = MONTO, NEGRITA
    ws.cell(r + 1, 1, len(conceptos) + 2)
    ws.cell(r + 1, 2, "Total Retenciones")
    for col in range(3, 9):
        ws.cell(r + 1, col, 0).number_format = MONTO
