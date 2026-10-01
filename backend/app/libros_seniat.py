"""Los libros de Compras y de Ventas en el formato del SENIAT, como .xlsx.

Copian las planillas que usa el contador: el detalle de cada documento,
los totales, las bases e IVA por alicuota, IGTF y retenciones, y el resumen
de creditos (compras) o debitos (ventas) fiscales al pie. Los montos van en
BOLIVARES: en compras, los congelados en cada factura al guardarla; en
ventas, a la tasa BCV congelada en cada venta al cobrarla.

Aca si se usa openpyxl, a diferencia de los CSV de `exportar_csv`: este
formato tiene dos filas de encabezado combinadas, sumas y un resumen debajo,
y un CSV no lo puede dar. Lo que el ERP todavia no lleva (importaciones,
IGTF, retenciones de IVA) sale en cero, como en las planillas.
"""

import datetime
import io
from typing import Callable, Dict, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import impuestos, schemas

GRIS = PatternFill("solid", fgColor="D9D9D9")
NEGRITA = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONTO = "#,##0.00"
PORCENTAJE = "0.00%"

DETALLE = [
    "N° operación", "Fecha del documento", "RIF", "Nombre/Razón social", "Tipo de Documento",
    "N° de Factura", "N° Nota de Crédito", "N° Nota de Débito", "N° de control", "Tipo de transacción",
    "N° Factura afectada",
]
POR_ALICUOTA = [
    "Base imponible (16%)", "Alicuota (16%)", "IVA 16%", "Base imponible (8%)", "Alicuota (8%)", "IVA 8%",
    "Base imponible (31%)", "Alicuota (31%)", "IVA 31%",
]

# ------------------------------------------------------------ compras
ENCABEZADOS = DETALLE + [
    "Total compras", "Total compras con IVA", "Total compras exentas",
] + POR_ALICUOTA + [
    "Total compras", "Total compras con IVA", "Total compras exentas",
    "Número de Declaración Única de Aduana", "Número de expediente de Importación",
    "Valor total de las importaciones definitivas",
    "Base imponible (16%)", "Alicuota Int. (16%)", "IVA Int. 16%", "Base imponible (8%)",
    "Alicuota Int. (8%)", "IVA Int. 8%", "Base imponible (31%)", "Alicuota Int. (31%)", "IVA Int. 31%",
    "Igtf", "Fecha Retención", "N° Retención", "IVA retenido",
]
GRUPOS_COMPRAS = [
    (1, 11, "DETALLE DEL DOCUMENTO"), (12, 14, "TOTALES"), (15, 23, "COMPRAS NACIONALES"),
    (24, 26, "TOTALES INTERNACIONALES"), (27, 38, "COMPRAS INTERNACIONALES"), (39, 39, "IGTF"),
    (40, 42, "RETENCIONES"),
]

# ------------------------------------------------------------ ventas
ENCABEZADOS_VENTAS = DETALLE + [
    "Total ventas", "Total ventas con IVA", "Total ventas exentas",
] + POR_ALICUOTA + ["Igtf", "Fecha Retención", "N° Retención", "IVA retenido"]
GRUPOS_VENTAS = [
    (1, 11, "DETALLE DEL DOCUMENTO"), (12, 14, "TOTALES"), (15, 17, "ALÍCUOTA GENERAL (16%)"),
    (18, 20, "ALÍCUOTA REDUCIDA (8%)"), (21, 23, "ALÍCUOTA ADICIONAL (31%)"), (24, 24, "IGTF"),
    (25, 27, "RETENCIONES"),
]

# Donde cae cada alicuota nacional, igual en los dos libros: (base, iva).
COLUMNAS_ALICUOTA = {16.0: (15, 17), 8.0: (18, 20), 31.0: (21, 23)}

FILA_GRUPOS = 6
FILA_ENCABEZADO = 7
PRIMERA_FILA = 8


def _f(x: datetime.date) -> str:
    return x.strftime("%d/%m/%Y")


def _cabecera(ws, fiscal, titulo: str, desde: datetime.date, hasta: datetime.date,
              grupos, encabezados: Sequence[str]) -> None:
    rif = impuestos.normalizar_rif(fiscal.rif or "")
    razon = (fiscal.razon_social or "").strip()
    ws["C1"] = " - ".join(x for x in (razon, rif) if x) or "Razón social y RIF: cárgalos en Impuestos"
    ws["C2"] = f"Direccion:  {fiscal.direccion}" if fiscal.direccion else "Direccion:"
    ws["C3"] = titulo
    ws["C4"] = f"Desde {_f(desde)} Hasta {_f(hasta)}"
    for fila in range(1, 5):
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=13)
        ws.cell(fila, 3).font = NEGRITA

    for desde_col, hasta_col, texto in grupos:
        if hasta_col > desde_col:
            ws.merge_cells(start_row=FILA_GRUPOS, start_column=desde_col, end_row=FILA_GRUPOS, end_column=hasta_col)
        for col in range(desde_col, hasta_col + 1):
            ws.cell(FILA_GRUPOS, col).fill = GRIS
        c = ws.cell(FILA_GRUPOS, desde_col, texto)
        c.font, c.alignment = NEGRITA, CENTRO
    for col, texto in enumerate(encabezados, start=1):
        c = ws.cell(FILA_ENCABEZADO, col, texto)
        c.font, c.fill, c.alignment = NEGRITA, GRIS, CENTRO


def _filas(ws, filas, montos: Sequence[int], alicuotas: Dict[int, int], alicuota: float,
           detalle: Callable[[int, object], Dict[int, object]]) -> int:
    """Escribe las filas desde PRIMERA_FILA y la de sumas. Devuelve la fila
    de sumas. Cada fila trae gravado_bs, exento_bs, iva_bs y total_bs (None
    si no tiene tasa)."""
    base_col, iva_col = COLUMNAS_ALICUOTA.get(alicuota, COLUMNAS_ALICUOTA[16.0])
    r = PRIMERA_FILA
    for n, f in enumerate(filas, start=1):
        for col, v in detalle(n, f).items():
            ws.cell(r, col, v)
        # Las planillas quieren los montos que no aplican en cero, no vacios.
        for col in montos:
            ws.cell(r, col, 0).number_format = MONTO
        for col, pct in alicuotas.items():
            ws.cell(r, col, pct / 100).number_format = PORCENTAJE
        if f.total_bs is not None:
            gravado, exento, iva = f.gravado_bs or 0, f.exento_bs or 0, f.iva_bs or 0
            ws.cell(r, 12, f.total_bs)
            ws.cell(r, 13, round(gravado + iva, 2))
            ws.cell(r, 14, exento)
            ws.cell(r, base_col, gravado)
            ws.cell(r, iva_col, iva)
        else:
            # Sin tasa no hay Bs: un cero pareceria un documento por nada.
            for col in (12, 13, 14, base_col, iva_col):
                ws.cell(r, col).value = None
            ws.cell(r, 4, f"{ws.cell(r, 4).value} (SIN TASA: falta la tasa de su fecha)")
        r += 1

    ultima = r - 1
    for col in montos:
        letra = get_column_letter(col)
        c = ws.cell(r, col, f"=SUM({letra}{PRIMERA_FILA}:{letra}{ultima})" if filas else 0)
        c.number_format, c.font = MONTO, NEGRITA
    return r


def _resumen(ws, r0: int, filas, alicuota: float, fiscales: str,
             conceptos: List[Tuple[str, Tuple[Optional[str], Optional[str]]]], total: str) -> None:
    """El cuadro del pie, con facturas y notas de credito por separado, como
    las planillas."""

    def suma(tipo: str, campo) -> float:
        if campo is None:
            return 0.0
        return round(sum(getattr(f, campo) or 0 for f in filas if f.tipo == tipo), 2)

    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=2)
    for col, texto in ((1, "Resumen"), (3, "Facturas/Notas de Débito"), (5, "Notas de Crédito"), (7, "Total Neto")):
        if col > 1:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r0, end_column=col + 1)
        c = ws.cell(r0, col, texto)
        c.font, c.fill, c.alignment = NEGRITA, GRIS, CENTRO
    credito = "Crédito Fiscal" if fiscales == "Créditos Fiscales" else "Débito Fiscal"
    subtitulos = [fiscales, "Base Imponible", credito, "Base Imponible", credito, "Base Imponible", credito]
    for col, texto in enumerate(subtitulos, start=2):
        c = ws.cell(r0 + 1, col, texto)
        c.font, c.fill, c.alignment = NEGRITA, GRIS, CENTRO

    r = r0 + 2
    for n, (texto, (base, impuesto)) in enumerate(conceptos, start=1):
        ws.cell(r, 1, n)
        ws.cell(r, 2, texto)
        montos = (suma("FAC", base), suma("FAC", impuesto), suma("NC", base), suma("NC", impuesto))
        for col, v in zip(range(3, 7), montos):
            ws.cell(r, col, v).number_format = MONTO
        ws.cell(r, 7, f"=C{r}+E{r}").number_format = MONTO
        ws.cell(r, 8, f"=D{r}+F{r}").number_format = MONTO
        r += 1
    ws.cell(r, 1, len(conceptos) + 1)
    ws.cell(r, 2, total).font = NEGRITA
    for col in range(3, 9):
        letra = get_column_letter(col)
        c = ws.cell(r, col, f"=SUM({letra}{r0 + 2}:{letra}{r - 1})")
        c.number_format, c.font = MONTO, NEGRITA
    ws.cell(r + 1, 1, len(conceptos) + 2)
    ws.cell(r + 1, 2, "Total Retenciones")
    for col in range(3, 9):
        ws.cell(r + 1, col, 0).number_format = MONTO


def _gravadas(alicuota: float, pct: float):
    return ("gravado_bs", "iva_bs") if pct == alicuota else (None, None)


def _terminar(ws, columnas: int) -> bytes:
    anchos = {1: 10, 2: 12, 3: 13, 4: 34, 5: 10, 6: 16, 7: 12, 8: 12, 9: 14, 10: 12, 11: 14}
    for col in range(1, columnas + 1):
        ws.column_dimensions[get_column_letter(col)].width = anchos.get(col, 15)
    ws.row_dimensions[FILA_ENCABEZADO].height = 45
    ws.freeze_panes = ws.cell(PRIMERA_FILA, 5)
    salida = io.BytesIO()
    ws.parent.save(salida)
    return salida.getvalue()


def libro_compras(libro: schemas.LibroCompras, fiscal, desde: datetime.date, hasta: datetime.date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Libro de Compras"
    _cabecera(ws, fiscal, "Libro de Compras", desde, hasta, GRUPOS_COMPRAS, ENCABEZADOS)

    def detalle(n, f):
        return {
            1: n, 2: _f(f.fecha_emision), 3: impuestos.normalizar_rif(f.proveedor_rif or ""),
            4: f.proveedor_nombre, 5: f.tipo,
            6: f.numero_factura if f.tipo == "FAC" else "",
            7: f.numero_nota if f.tipo == "NC" else "", 8: "",
            9: f.numero_control, 10: "01-REG", 11: f.factura_afectada or "--",
            27: "-", 28: "--",
        }

    alicuota = float(libro.tasa_iva)
    montos = [12, 13, 14, 15, 17, 18, 20, 21, 23, 24, 25, 26, 29, 30, 32, 33, 35, 36, 38, 39, 42]
    alicuotas = {16: 16, 19: 8, 22: 31, 31: 16, 34: 8, 37: 31}
    r = _filas(ws, libro.filas, montos, alicuotas, alicuota, detalle)
    _resumen(ws, r + 2, libro.filas, alicuota, "Créditos Fiscales", [
        ("Compras Internas no Gravadas", ("exento_bs", None)),
        ("Importaciones Gravadas por Alícuota Reducida", (None, None)),
        ("Importaciones Gravadas por Alícuota General", (None, None)),
        ("Importaciones Gravadas por Alícuota General más Adicional", (None, None)),
        ("Compras Internas Gravadas sólo por Alícuota General", _gravadas(alicuota, 16.0)),
        ("Compras Internas Gravadas por Alícuota General más Adicional", _gravadas(alicuota, 31.0)),
        ("Compras Internas Gravadas por Alícuota Reducida", _gravadas(alicuota, 8.0)),
        ("Ajustes a los Créditos Fiscales de Periodos Anteriores", (None, None)),
    ], "Total Compras y Créditos Fiscales del Periodo")
    return _terminar(ws, len(ENCABEZADOS))


def libro_ventas(libro: schemas.LibroVentas, fiscal, desde: datetime.date, hasta: datetime.date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Libro de Ventas"
    _cabecera(ws, fiscal, "Libro de Ventas", desde, hasta, GRUPOS_VENTAS, ENCABEZADOS_VENTAS)

    def detalle(n, f):
        return {
            1: n, 2: _f(f.fecha.date()), 3: impuestos.normalizar_rif(f.rif or ""), 4: f.cliente,
            5: f.tipo,
            6: f.numero_factura if f.tipo == "FAC" else "--",
            7: f.numero_nota if f.tipo == "NC" else "--", 8: "--",
            9: f.numero_control, 10: "01-REG", 11: f.factura_afectada or "--",
        }

    alicuota = float(libro.tasa_iva)
    montos = [12, 13, 14, 15, 17, 18, 20, 21, 23, 24, 27]
    alicuotas = {16: 16, 19: 8, 22: 31}
    r = _filas(ws, libro.filas, montos, alicuotas, alicuota, detalle)
    _resumen(ws, r + 2, libro.filas, alicuota, "Débitos Fiscales", [
        ("Ventas Internas no Gravadas", ("exento_bs", None)),
        ("Ventas de Exportación", (None, None)),
        ("Ventas Internas Gravadas sólo por Alícuota General", _gravadas(alicuota, 16.0)),
        ("Ventas Internas Gravadas por Alícuota Reducida", _gravadas(alicuota, 8.0)),
        ("Ventas Internas Gravadas por Alícuota General más Adicional", _gravadas(alicuota, 31.0)),
        ("Ajustes a los Débitos Fiscales de Periodos Anteriores", (None, None)),
    ], "Total Ventas y Débitos Fiscales del Periodo")
    return _terminar(ws, len(ENCABEZADOS_VENTAS))
