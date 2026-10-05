"""Las retenciones de IVA que le hacen a la empresa sus clientes.

Siendo contribuyente ordinario, cuando le vende a un contribuyente especial
este le retiene el 75 % del IVA y le entrega un comprobante. Ese monto va en
el Libro de Ventas (Y-AA) del mes del comprobante y se descuenta del IVA a
pagar en la declaracion; lo que sobra se arrastra.
"""

import datetime
import io

import pytest
from openpyxl import load_workbook

from app import impuestos, models
from app.timeutils import hoy

TASA = 100.0


@pytest.fixture(autouse=True)
def tasa(db):
    db.merge(models.TasaCambio(fecha=datetime.date(2000, 1, 1), bcv=TASA, origen="auto"))
    db.commit()


def vender(client, variante, numero="V-1"):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 10}], "nota": ""}).json()
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={
        "metodo_pago": "Efectivo Bs", "facturado": True, "numero_factura": numero,
        "rif_cliente": "J300215938", "razon_social_cliente": "Diageo Venezuela, C.A.",
    })
    assert r.status_code == 200, r.text
    return r.json()


def comprobante(fecha):
    return f"{fecha.year:04d}{fecha.month:02d}00001234"


def iva_bs(db, pedido_id):
    p = db.get(models.Pedido, pedido_id)
    return impuestos.desglosar(round(p.total * TASA, 2), 16)[1]


def test_por_defecto_se_registra_el_75_del_iva(client, db, variante):
    v = vender(client, variante)
    r = client.post(f"/api/impuestos/ventas/{v['id']}/retencion", json={
        "comprobante": comprobante(hoy()), "fecha": hoy().isoformat(),
    })
    assert r.status_code == 200, r.text
    assert r.json()["iva_retenido_bs"] == impuestos.porcentaje_de(iva_bs(db, v["id"]), 75)

    [f] = client.get("/api/impuestos/libro-ventas").json()["filas"]
    assert (f["comprobante_retencion"], f["iva_retenido_bs"]) == (comprobante(hoy()), r.json()["iva_retenido_bs"])


def test_el_comprobante_se_valida(client, variante):
    v = vender(client, variante)
    url = f"/api/impuestos/ventas/{v['id']}/retencion"
    assert client.post(url, json={"comprobante": "123", "fecha": hoy().isoformat()}).status_code == 400
    manana = (hoy() + datetime.timedelta(days=1)).isoformat()
    assert client.post(url, json={"comprobante": comprobante(hoy()), "fecha": manana}).status_code == 400
    assert client.post(url, json={"comprobante": comprobante(hoy()), "fecha": hoy().isoformat(), "monto_bs": 10**9}).status_code == 400


def test_la_declaracion_descuenta_lo_retenido(client, db, variante):
    v = vender(client, variante)
    # Venta y comprobante del mes pasado, para poder declararlo.
    p = db.get(models.Pedido, v["id"])
    p.cerrado_en = p.cerrado_en.replace(day=1) - datetime.timedelta(days=20)
    db.commit()
    fecha = p.cerrado_en.date()
    client.post(f"/api/impuestos/ventas/{v['id']}/retencion", json={"comprobante": comprobante(fecha), "fecha": fecha.isoformat()})

    d = client.post("/api/impuestos/declaraciones", json={"anio": fecha.year, "mes": fecha.month}).json()
    iva = iva_bs(db, v["id"])
    retenido = impuestos.porcentaje_de(iva, 75)
    assert d["iva_debito_bs"] == iva
    assert d["retenciones_usadas_bs"] == retenido
    assert d["iva_a_pagar_bs"] == round(iva - retenido, 2)


def test_un_comprobante_de_otro_mes_va_en_el_libro_de_ese_mes(client, db, variante):
    v = vender(client, variante)
    p = db.get(models.Pedido, v["id"])
    p.cerrado_en = p.cerrado_en.replace(day=1) - datetime.timedelta(days=20)
    db.commit()
    client.post(f"/api/impuestos/ventas/{v['id']}/retencion", json={
        "comprobante": comprobante(hoy()), "fecha": hoy().isoformat(),
    })
    filas = client.get("/api/impuestos/libro-ventas").json()["filas"]
    [ret] = [f for f in filas if f["tipo"] == "RET"]
    assert ret["factura_afectada"] == "V-1" and ret["iva_retenido_bs"] > 0 and ret["total_bs"] == 0

    ws = load_workbook(io.BytesIO(client.get("/api/impuestos/libro-ventas/seniat").content)).active
    assert (ws["Z8"].value, ws["AA8"].value) == (comprobante(hoy()), ret["iva_retenido_bs"])


def test_se_puede_quitar_si_el_mes_no_esta_declarado(client, variante):
    v = vender(client, variante)
    url = f"/api/impuestos/ventas/{v['id']}/retencion"
    client.post(url, json={"comprobante": comprobante(hoy()), "fecha": hoy().isoformat()})
    assert client.delete(url).status_code == 200
    [f] = client.get("/api/impuestos/libro-ventas").json()["filas"]
    assert f["iva_retenido_bs"] is None


# ------------------------------------------------------- retenida en la caja

def cobrar_con_retencion(client, variante, pct=75, **extra):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 10}], "nota": ""}).json()
    return p, client.post(f"/api/pedidos/{p['id']}/cobrar", json={
        "metodo_pago": "Efectivo Bs", "facturado": True, "numero_factura": "V-9",
        "rif_cliente": "J300215938", "retencion_iva_pct": pct, **extra,
    })


def test_en_caja_se_cobra_el_total_menos_lo_retenido(client, db, variante):
    from conftest import saldo
    p, r = cobrar_con_retencion(client, variante)
    assert r.status_code == 200, r.text
    v = r.json()
    iva = impuestos.desglosar(v["total"], 16)[1]
    retenido = impuestos.porcentaje_de(iva, 75)
    assert v["retencion_iva_usd"] == retenido
    pagos = {pg["metodo"]: pg["monto"] for pg in v["pagos"]}
    assert pagos == {"Efectivo Bs": round(v["total"] - retenido, 2), "Retención IVA": retenido}
    # A la gaveta entro solo lo cobrado; lo retenido queda en 1035.
    assert saldo(db, "1035") == retenido
    # Queda esperando el comprobante, con su monto en Bs.
    [f] = client.get("/api/impuestos/libro-ventas").json()["filas"]
    assert f["retencion_pendiente_bs"] == impuestos.porcentaje_de(iva_bs(db, p["id"]), 75) and f["iva_retenido_bs"] is None


def test_los_pagos_tienen_que_sumar_lo_que_queda(client, variante):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 10}], "nota": ""}).json()
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={
        "metodo_pago": "Efectivo Bs", "facturado": True, "numero_factura": "V-8", "retencion_iva_pct": 75,
        "pagos": [{"metodo": "Efectivo Bs", "monto": p["total"]}],
    })
    assert r.status_code == 400 and "hay que cobrar" in r.json()["detail"]


def test_sin_factura_no_se_retiene_y_no_se_elige_como_pago(client, variante):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 1}], "nota": ""}).json()
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={"metodo_pago": "Efectivo Bs", "retencion_iva_pct": 75})
    assert r.status_code == 400
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={"metodo_pago": "Retención IVA"})
    assert r.status_code == 400 and "no se elige" in r.json()["detail"]


def test_el_comprobante_toma_el_monto_retenido_y_la_declaracion_salda_1035(client, db, variante):
    from conftest import saldo
    p, r = cobrar_con_retencion(client, variante)
    v = r.json()
    ped = db.get(models.Pedido, p["id"])
    ped.cerrado_en = ped.cerrado_en.replace(day=1) - datetime.timedelta(days=20)
    db.commit()
    fecha = ped.cerrado_en.date()
    c = client.post(f"/api/impuestos/ventas/{p['id']}/retencion", json={"comprobante": comprobante(fecha), "fecha": fecha.isoformat()})
    assert c.status_code == 200, c.text
    assert c.json()["iva_retenido_bs"] == ped.retencion_iva_bs

    d = client.post("/api/impuestos/declaraciones", json={"anio": fecha.year, "mes": fecha.month}).json()
    assert d["retenciones_usadas"] == v["retencion_iva_usd"]
    assert d["retenciones_usadas_bs"] == ped.retencion_iva_bs
    assert saldo(db, "1035") == 0
