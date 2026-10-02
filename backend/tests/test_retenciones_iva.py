"""Retenciones de IVA siendo agente de retencion (contribuyente especial).

Al proveedor se le paga el total menos lo retenido; lo retenido se le debe al
SENIAT (2050) hasta enterarlo por quincena. Cada retencion lleva su
comprobante (AAAAMM + correlativo de 8) y sale en el TXT de la quincena:
16 campos separados por tabulador, montos en Bs.
"""

import datetime
import io

import pytest
from openpyxl import load_workbook

from app import models
from app.timeutils import hoy

from conftest import saldo

TASA = 50.0
RIF_EMPRESA = "J402235135"


@pytest.fixture()
def agente(client, db):
    db.merge(models.TasaCambio(fecha=datetime.date(2000, 1, 1), bcv=TASA, origen="auto"))
    db.commit()
    client.put("/api/impuestos/config", json={
        "tasa_iva": 16, "razon_social": "Savora, C.A.", "rif": RIF_EMPRESA, "agente_retencion": True,
    })


def comprar(client, insumo, numero="F-1", costo=100.0, forma_pago="Efectivo", **extra):
    r = client.post("/api/compras/facturas", json={
        "numero_factura": numero, "numero_control": f"00-{numero}", "proveedor_nombre": "Proveedor",
        "proveedor_rif": "J-30021593-8", "forma_pago": forma_pago,
        "items": [{"ingrediente_id": insumo.id, "cantidad": 1, "costo_unitario": costo, "exento": False}],
        **extra,
    })
    assert r.status_code == 200, r.text
    return r.json()


def quincena_de_hoy():
    d = hoy()
    return d.year, d.month, 1 if d.day <= 15 else 2


def test_sin_ser_agente_no_se_retiene(client, insumo):
    f = comprar(client, insumo)
    assert (f["iva_retenido"], f["comprobante_retencion"], f["a_pagar"]) == (0, "", f["total"])


def test_se_retiene_el_75_y_al_proveedor_se_le_paga_el_resto(client, db, insumo, agente):
    f = comprar(client, insumo, forma_pago="Credito")
    assert (f["retencion_pct"], f["iva_retenido"], f["iva_retenido_bs"]) == (75, 12.0, 600.0)
    assert f["a_pagar"] == 104.0
    hoy_ = hoy()
    assert f["comprobante_retencion"] == f"{hoy_.year:04d}{hoy_.month:02d}00000001"
    # Se le debe al proveedor 104 y al SENIAT 12.
    assert (saldo(db, "2010"), saldo(db, "2050")) == (104.0, 12.0)
    # Pagarle al proveedor salda solo lo suyo.
    client.post(f"/api/compras/facturas/{f['id']}/pagar", json={"forma_pago": "Banco", "referencia": "123"})
    assert saldo(db, "2010") == 0


def test_el_100_y_el_correlativo_sigue(client, db, insumo, agente):
    comprar(client, insumo, "F-1")
    f = comprar(client, insumo, "F-2", retencion_pct=100)
    assert (f["iva_retenido"], f["comprobante_retencion"][-8:]) == (16.0, "00000002")
    # El proveedor queda recordado al 100 % para la proxima.
    client.post("/api/proveedores", json={"nombre": "Proveedor", "rif": "J300215938"})
    comprar(client, insumo, "F-3", retencion_pct=100)
    assert comprar(client, insumo, "F-4")["retencion_pct"] == 100


def test_un_porcentaje_que_no_existe_se_rechaza(client, insumo, agente):
    r = client.post("/api/compras/facturas", json={
        "numero_factura": "F-9", "proveedor_nombre": "P", "proveedor_rif": "J300215938", "retencion_pct": 50,
        "items": [{"ingrediente_id": insumo.id, "cantidad": 1, "costo_unitario": 1.0}],
    })
    assert r.status_code == 400


def test_el_txt_de_la_quincena(client, insumo, agente):
    comprar(client, insumo, "F-1")
    anio, mes, q = quincena_de_hoy()
    r = client.get(f"/api/impuestos/retenciones-iva/txt?anio={anio}&mes={mes}&quincena={q}")
    assert r.status_code == 200, r.text
    [linea] = r.content.decode().strip().split("\r\n")
    campos = linea.split("\t")
    assert len(campos) == 16
    assert campos[:8] == [RIF_EMPRESA, f"{anio:04d}{mes:02d}", hoy().isoformat(), "C", "01", "J300215938", "F-1", "00-F-1"]
    # Total, base y retenido en Bs; sin documento afectado; alicuota; sin expediente.
    assert campos[8:] == ["5800.00", "5000.00", "600.00", "0", f"{anio:04d}{mes:02d}00000001", "0.00", "16.00", "0"]


def test_una_quincena_sin_retenciones_da_un_txt_vacio(client, agente):
    anio, mes, q = quincena_de_hoy()
    r = client.get(f"/api/impuestos/retenciones-iva/txt?anio={anio}&mes={mes}&quincena={q}")
    assert r.status_code == 200 and r.content == b""


def test_sin_rif_de_la_empresa_no_hay_txt(client, db, insumo, agente):
    client.put("/api/impuestos/config", json={"tasa_iva": 16, "rif": ""})
    anio, mes, q = quincena_de_hoy()
    assert client.get(f"/api/impuestos/retenciones-iva/txt?anio={anio}&mes={mes}&quincena={q}").status_code == 400


def test_enterar_la_quincena_salda_la_deuda_con_el_seniat(client, db, insumo, agente):
    comprar(client, insumo, "F-1")
    anio, mes, q = quincena_de_hoy()
    r = client.post("/api/impuestos/retenciones-iva/enterar", json={"anio": anio, "mes": mes, "quincena": q, "forma_pago": "Banco"})
    assert r.status_code == 200, r.text
    assert r.json()["enterada"] is True and r.json()["total_retenido_bs"] == 600.0
    assert saldo(db, "2050") == 0
    otra = client.post("/api/impuestos/retenciones-iva/enterar", json={"anio": anio, "mes": mes, "quincena": q})
    assert otra.status_code == 409


def test_el_libro_de_compras_lleva_la_retencion(client, insumo, agente):
    comprar(client, insumo, "F-1")
    [fila] = client.get("/api/impuestos/libro-compras").json()["filas"]
    assert (fila["iva_retenido_bs"], fila["comprobante_retencion"][-8:]) == (600.0, "00000001")
    ws = load_workbook(io.BytesIO(client.get("/api/impuestos/libro-compras/seniat").content)).active
    assert (ws["AN8"].value, ws["AP8"].value) == (hoy().strftime("%d/%m/%Y"), 600.0)
    assert ws["AO8"].value.endswith("00000001")
