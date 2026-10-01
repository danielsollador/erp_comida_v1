"""El Libro de Compras en bolivares, con el formato del SENIAT.

El ERP lleva todo en dolares, pero el libro que se declara va en Bs. Lo que
importa fijar:
  * una factura en $ pasa a Bs a la tasa BCV de su fecha de EMISION, no la
    del dia en que se carga (o a la tasa que imprime el papel);
  * una factura en Bs conserva exactos los Bs del papel, aunque por dentro se
    haya guardado en dolares;
  * lo exento, lo gravado y el IVA van en columnas aparte;
  * las notas de credito son filas propias, en negativo;
  * el .xlsx tiene las columnas de la planilla y sus totales cuadran.
"""

import datetime
import io

import pytest
from openpyxl import load_workbook

from app import libro_compras_seniat, models
from app.timeutils import hoy

EMISION = hoy() - datetime.timedelta(days=8)
TASA_EMISION = 850.0
TASA_HOY = 900.0


@pytest.fixture()
def tasas(db):
    """La tasa del dia de la factura, y otra (mas alta) despues."""
    for fecha in (EMISION, EMISION - datetime.timedelta(days=1)):
        db.merge(models.TasaCambio(fecha=fecha, bcv=TASA_EMISION, origen="auto"))
    db.merge(models.TasaCambio(fecha=hoy() - datetime.timedelta(days=1), bcv=TASA_HOY, origen="auto"))
    db.commit()


@pytest.fixture()
def harina(db):
    ing = models.Ingrediente(nombre="Harina", unidad="kg", stock_actual=0, stock_minimo=0,
                             stock_objetivo=0, costo_unitario=1, exento=True)
    db.add(ing)
    db.commit()
    return ing


def cargar(client, insumo, **extra):
    cuerpo = {
        "numero_factura": "F-1", "proveedor_nombre": "Proveedor", "proveedor_rif": "J402235135",
        "categoria": "Insumos", "forma_pago": "Efectivo",
        "fecha_emision": EMISION.isoformat(),
        "items": [{"ingrediente_id": insumo.id, "cantidad": 2, "costo_unitario": 50.0}],
    }
    cuerpo.update(extra)
    r = client.post("/api/compras/facturas", json=cuerpo)
    assert r.status_code == 200, r.text
    return r.json()


def libro(client):
    return client.get("/api/impuestos/libro-compras").json()


def test_una_factura_en_dolares_va_a_la_tasa_de_su_fecha_de_emision(client, insumo, tasas):
    """La factura es de hace 8 dias y se carga hoy: vale en Bs lo que valia
    ese dia. Con la tasa de hoy el libro diria otro monto."""
    f = cargar(client, insumo, numero_control="00-0001")
    assert f["moneda"] == "$" and f["tasa_bcv"] == TASA_EMISION and f["numero_control"] == "00-0001"

    fila = libro(client)["filas"][0]
    assert fila["tipo"] == "FAC" and fila["numero_control"] == "00-0001"
    assert fila["gravado_bs"] == 100 * TASA_EMISION
    assert fila["iva_bs"] == round(100 * TASA_EMISION * 0.16, 2)
    assert fila["total_bs"] == round(100 * TASA_EMISION * 1.16, 2)
    assert fila["tasa_estimada"] is False


def test_la_tasa_que_imprime_el_papel_manda(client, insumo, tasas):
    """Diageo imprime su tasa (857,006): con esa calculo su IVA en Bs."""
    cargar(client, insumo, tasa_bcv=857.006)
    fila = libro(client)["filas"][0]
    assert fila["tasa_bcv"] == 857.006
    assert fila["gravado_bs"] == round(100 * 857.006, 2)


def test_una_factura_en_bolivares_conserva_los_bolivares_del_papel(client, insumo, harina, tasas):
    """El formulario pasa los Bs a dolares dividiendo entre la tasa; el libro
    tiene que devolver los Bs del papel al centimo, no la vuelta redondeada."""
    tasa = 861.37
    gravado_papel, exento_papel = 12370.73, 1289.99
    cargar(
        client, insumo, moneda="Bs", tasa_bcv=tasa,
        items=[
            {"ingrediente_id": insumo.id, "cantidad": 3, "costo_unitario": gravado_papel / 3 / tasa, "exento": False},
            {"ingrediente_id": harina.id, "cantidad": 1, "costo_unitario": exento_papel / tasa},
        ],
    )
    fila = libro(client)["filas"][0]
    assert fila["moneda"] == "Bs"
    assert fila["gravado_bs"] == gravado_papel
    assert fila["exento_bs"] == exento_papel
    assert fila["iva_bs"] == round(gravado_papel * 0.16, 2)
    assert fila["total_bs"] == round(gravado_papel + exento_papel + round(gravado_papel * 0.16, 2), 2)


def test_recargo_y_descuento_en_bolivares_tambien_son_exactos(client, insumo, tasas):
    tasa = 860.18
    cargar(
        client, insumo, moneda="Bs", tasa_bcv=tasa, recargo=1000 / tasa, descuento=250.5 / tasa,
        items=[{"ingrediente_id": insumo.id, "cantidad": 1, "costo_unitario": 10000 / tasa}],
    )
    assert libro(client)["filas"][0]["gravado_bs"] == 10749.5


def test_sin_tasa_una_factura_en_bolivares_no_se_guarda(client, insumo):
    r = client.post("/api/compras/facturas", json={
        "numero_factura": "F-1", "proveedor_nombre": "P", "proveedor_rif": "J402235135",
        "fecha_emision": "2020-01-01", "moneda": "Bs",
        "items": [{"ingrediente_id": insumo.id, "cantidad": 1, "costo_unitario": 1.0}],
    })
    assert r.status_code == 400 and "tasa" in r.json()["detail"]


def test_una_factura_en_dolares_sin_tasa_queda_sin_bolivares_y_se_avisa(client, insumo, db):
    """Una fecha de antes de que el ERP guardara tasas: hueco visible, no un
    monto inventado."""
    cargar(client, insumo, fecha_emision="2020-01-01")
    lib = libro(client)
    assert lib["filas"][0]["total_bs"] is None
    assert lib["sin_tasa"] == 1


def test_una_factura_de_antes_se_pasa_con_la_tasa_guardada_de_su_fecha(client, insumo, tasas, db):
    """Las facturas cargadas antes de este cambio no tienen Bs congelados."""
    cargar(client, insumo)
    f = db.query(models.FacturaCompra).one()
    f.tasa_bcv = f.gravado_bs = f.exento_bs = f.iva_bs = None
    db.commit()
    fila = libro(client)["filas"][0]
    assert fila["tasa_bcv"] == TASA_EMISION and fila["tasa_estimada"] is True
    assert fila["gravado_bs"] == 100 * TASA_EMISION


def test_la_nota_de_credito_es_su_propia_fila_en_negativo(client, insumo, tasas):
    f = cargar(client, insumo)
    r = client.post(f"/api/compras/facturas/{f['id']}/notas-credito", json={
        "numero": "NC-7", "tipo": "descuento", "base_imponible": 10.0,
    })
    assert r.status_code == 200, r.text
    fac, nc = libro(client)["filas"]
    assert (nc["tipo"], nc["numero_nota"], nc["factura_afectada"]) == ("NC", "NC-7", "F-1")
    assert nc["gravado_bs"] == -10 * TASA_EMISION
    assert nc["base_imponible"] == -10
    lib = libro(client)
    # En dolares, el total sigue siendo el neto de antes (lo que declara el IVA).
    assert lib["total_base"] == 90
    assert lib["total_gravado_bs"] == 90 * TASA_EMISION


def test_el_excel_tiene_el_formato_del_seniat(client, insumo, harina, tasas):
    client.put("/api/impuestos/config", json={
        "tasa_iva": 16, "razon_social": "Inversiones Savora, C.A.", "rif": "J-40223513-5",
        "direccion": "Av. Principal",
    })
    f = cargar(client, insumo, numero_control="00-0001", items=[
        {"ingrediente_id": insumo.id, "cantidad": 1, "costo_unitario": 100.0},
        {"ingrediente_id": harina.id, "cantidad": 1, "costo_unitario": 20.0},
    ])
    client.post(f"/api/compras/facturas/{f['id']}/notas-credito", json={
        "numero": "NC-7", "tipo": "descuento", "base_imponible": 10.0,
    })

    r = client.get("/api/impuestos/libro-compras/seniat")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active

    assert ws["C1"].value == "Inversiones Savora, C.A. - J402235135"
    assert ws["C3"].value == "Libro de Compras"
    assert [ws.cell(7, c).value for c in range(1, 43)] == libro_compras_seniat.ENCABEZADOS
    assert ws["A6"].value == "DETALLE DEL DOCUMENTO" and ws["AN6"].value == "RETENCIONES"

    fila = {ws.cell(7, c).value + str(c): ws.cell(8, c).value for c in range(1, 43)}
    assert fila["Tipo de Documento5"] == "FAC" and fila["N° de control9"] == "00-0001"
    assert fila["Fecha del documento2"] == EMISION.strftime("%d/%m/%Y")
    assert fila["RIF3"] == "J402235135"
    gravado, exento = 100 * TASA_EMISION, 20 * TASA_EMISION
    iva = round(gravado * 0.16, 2)
    assert (ws["L8"].value, ws["M8"].value, ws["N8"].value) == (round(gravado + exento + iva, 2), round(gravado + iva, 2), exento)
    assert (ws["O8"].value, ws["P8"].value, ws["Q8"].value) == (gravado, 0.16, iva)
    # La nota de credito, debajo, en negativo y apuntando a su factura.
    assert (ws["E9"].value, ws["G9"].value, ws["K9"].value) == ("NC", "NC-7", "F-1")
    assert ws["L9"].value < 0
    # Totales al pie.
    assert ws["L10"].value == "=SUM(L8:L9)"


def test_cambiar_la_alicuota_no_borra_la_razon_social(client):
    client.put("/api/impuestos/config", json={"tasa_iva": 16, "razon_social": "Savora", "rif": "J402235135"})
    client.put("/api/impuestos/config", json={"tasa_iva": 16})
    c = client.get("/api/impuestos/config").json()
    assert (c["razon_social"], c["rif"]) == ("Savora", "J402235135")


def test_la_tasa_de_una_fecha(client, tasas):
    sabado = EMISION + datetime.timedelta(days=1)
    r = client.get(f"/api/tasas/al?fecha={sabado.isoformat()}").json()
    assert r["bcv"] == TASA_EMISION and r["fecha"] == EMISION.isoformat()
    assert client.get("/api/tasas/al?fecha=2001-01-01").json()["bcv"] is None
