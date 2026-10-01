"""El Libro de Ventas en bolivares, con el formato del SENIAT.

Cada venta facturada congela la tasa BCV al cobrarla: con esa pasa a Bs, no
con la de hoy. El IVA se desglosa sobre los Bs, como lo imprime la maquina
fiscal. La nota de credito de una devolucion es su propia fila, en negativo,
que apunta a la factura que afecta.
"""

import io

from openpyxl import load_workbook

from app import impuestos, libros_seniat, models

from test_devolucion_periodo_declarado import envejecer, libro, mes_de, vendida_y_facturada

TASA = 850.0


def congelar_tasa(db, pedido_id, tasa=TASA):
    p = db.get(models.Pedido, pedido_id)
    p.tasa_bcv = tasa
    db.commit()
    return p


def test_una_venta_pasa_a_bs_con_la_tasa_de_su_cobro(client, db, variante):
    v = vendida_y_facturada(client, variante, "F-1")
    p = congelar_tasa(db, v["id"])
    inicio, fin = mes_de(p.cerrado_en)

    [f] = libro(client, inicio, fin)["filas"]
    total_bs = round(p.total * TASA, 2)
    base_bs, iva_bs = impuestos.desglosar(total_bs, p.tasa_iva or 16)
    assert (f["tipo"], f["tasa_bcv"], f["total_bs"]) == ("FAC", TASA, total_bs)
    assert (f["gravado_bs"], f["iva_bs"], f["exento_bs"]) == (base_bs, iva_bs, 0)


def test_la_nota_de_credito_es_su_propia_fila_que_apunta_a_la_factura(client, db, variante):
    v = vendida_y_facturada(client, variante, "F-2")
    congelar_tasa(db, v["id"])
    envejecer(db, v["id"], dias=40)
    r = client.post(f"/api/pedidos/{v['id']}/devolver", json={
        "recuperable": True, "motivo": "no le gusto", "nota_credito": "NC-2",
    })
    assert r.status_code == 200, r.text

    p = db.get(models.Pedido, v["id"])
    inicio, fin = mes_de(p.fecha_devolucion)
    [nc] = libro(client, inicio, fin)["filas"]
    assert (nc["tipo"], nc["numero_nota"], nc["factura_afectada"]) == ("NC", "NC-2", "F-2")
    assert nc["total_bs"] == -round(p.total * TASA, 2)


def test_el_excel_tiene_el_formato_del_seniat(client, db, variante):
    client.put("/api/impuestos/config", json={
        "tasa_iva": 16, "razon_social": "Inversiones Savora, C.A.", "rif": "J402235135",
    })
    v = vendida_y_facturada(client, variante, "F-3")
    p = congelar_tasa(db, v["id"])
    inicio, fin = mes_de(p.cerrado_en)
    hasta = (fin - __import__("datetime").timedelta(days=1)).date()

    r = client.get(f"/api/impuestos/libro-ventas/seniat?desde={inicio.date()}&hasta={hasta}")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active

    assert ws["C1"].value == "Inversiones Savora, C.A. - J402235135"
    assert ws["C3"].value == "Libro de Ventas"
    assert [ws.cell(7, c).value for c in range(1, 28)] == libros_seniat.ENCABEZADOS_VENTAS
    assert ws["O6"].value == "ALÍCUOTA GENERAL (16%)" and ws["Y6"].value == "RETENCIONES"

    total_bs = round(p.total * TASA, 2)
    base_bs, iva_bs = impuestos.desglosar(total_bs, 16)
    assert (ws["E8"].value, ws["F8"].value, ws["G8"].value, ws["J8"].value) == ("FAC", "F-3", "--", "01-REG")
    assert ws["D8"].value == "Consumidor final"
    assert (ws["L8"].value, ws["M8"].value, ws["N8"].value) == (total_bs, round(base_bs + iva_bs, 2), 0)
    assert (ws["O8"].value, ws["P8"].value, ws["Q8"].value) == (base_bs, 0.16, iva_bs)
    assert ws["L9"].value == "=SUM(L8:L8)"


def test_al_facturar_se_guardan_el_rif_y_el_control_y_van_al_libro(client, db, variante):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 1}], "nota": ""}).json()
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={
        "metodo_pago": "Efectivo Bs", "facturado": True, "numero_factura": "F-9",
        "numero_control": "00-000123", "rif_cliente": "j-40223513-5", "razon_social_cliente": "Megalicores, C.A.",
    })
    assert r.status_code == 200, r.text
    assert (r.json()["rif_cliente"], r.json()["numero_control"]) == ("J402235135", "00-000123")

    pedido = db.get(models.Pedido, p["id"])
    inicio, fin = mes_de(pedido.cerrado_en)
    [f] = libro(client, inicio, fin)["filas"]
    assert (f["rif"], f["cliente"], f["numero_control"]) == ("J402235135", "Megalicores, C.A.", "00-000123")


def test_una_cedula_sola_se_lee_como_venezolana(client, variante):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 1}], "nota": ""}).json()
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={
        "metodo_pago": "Efectivo Bs", "facturado": True, "numero_factura": "F-10", "rif_cliente": "12.345.678",
    })
    assert r.json()["rif_cliente"] == "V12345678"


def test_un_rif_mal_escrito_no_se_factura(client, variante):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 1}], "nota": ""}).json()
    r = client.post(f"/api/pedidos/{p['id']}/cobrar", json={
        "metodo_pago": "Efectivo Bs", "facturado": True, "numero_factura": "F-11", "rif_cliente": "sin rif",
    })
    assert r.status_code == 400 and "cédula" in r.json()["detail"]


def test_facturar_despues_tambien_pide_rif_y_control(client, variante):
    p = client.post("/api/pedidos", json={"items": [{"variante_id": variante.id, "cantidad": 1}], "nota": ""}).json()
    client.post(f"/api/pedidos/{p['id']}/cobrar", json={"metodo_pago": "Efectivo Bs"})
    r = client.post(f"/api/pedidos/{p['id']}/facturar", json={
        "numero_factura": "F-12", "numero_control": "00-9", "rif_cliente": "V12345678", "razon_social_cliente": "Ana Pérez",
    })
    assert r.status_code == 200, r.text
    assert (r.json()["rif_cliente"], r.json()["razon_social_cliente"]) == ("V12345678", "Ana Pérez")
